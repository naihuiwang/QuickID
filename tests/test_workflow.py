"""Regression checks for shuffled headings, marker matching, and CSV output."""
import csv
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from decimal import Decimal
from quickid_core import HEADINGS, default_settings, match, process, read_peaks, read_reference
from quickid import checked_settings, load_settings, save_settings, stop_remembering_settings
from quickid import App, MARKERS as GUI_MARKERS
from unittest.mock import Mock, patch


class Workflow(unittest.TestCase):
    def test_peak_masses_round_half_up_to_one_decimal_before_matching(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'peaks.txt'
            path.write_text('m/z\tintensity\n1200.04\t1\n1200.05\t1\n'
                            '1200.149\t1\n1200.15\t1\n1200.250\t1\n', encoding='utf-8')
            self.assertEqual(read_peaks(path), [Decimal(value) for value in
                             ('1200.0', '1200.1', '1200.1', '1200.2', '1200.3')])
            row = dict.fromkeys(HEADINGS, '')
            row['B'] = '1200.1'
            _, hits, _ = match([row], read_peaks(path), default_settings())
            self.assertEqual(hits[3], '1200.0/1200.1/1200.2/1200.3')

    def test_peak_count_warning_is_optional_and_uses_editable_threshold(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            book = Workbook()
            book.active.append(HEADINGS)
            book.active.append(['', '', 'Taxon', *([''] * 12)])
            book.save(root / 'ref.xlsx')
            for amount in (500, 501):
                (root / f'peaks_{amount}.txt').write_text(
                    'm/z\tintensity\n' + '900.1\t1\n' * amount, encoding='utf-8')
            output = root / 'results.csv'
            def summaries_for(settings):
                process(root, output, settings, reference_file=root / 'ref.xlsx')
                with output.open(encoding='utf-8-sig', newline='') as stream:
                    records = list(csv.reader(stream))
                self.assertEqual(records[0][-1], 'QuickID mismatch flags')
                return {row[0]: row for row in records[1:] if row and row[0]}

            defaults = default_settings()
            self.assertEqual(defaults['peak_warning_enabled'], '1')
            self.assertEqual(defaults['peak_warning_threshold'], '500')
            summaries = summaries_for(defaults)
            self.assertEqual(summaries['peaks_500.txt'][-1], '')
            self.assertEqual(summaries['peaks_501.txt'][-1],
                             'Possible false positive (>500 peaks)')
            summaries = summaries_for(checked_settings({**defaults, 'peak_warning_threshold': '499'}))
            self.assertEqual(summaries['peaks_500.txt'][-1],
                             'Possible false positive (>499 peaks)')
            summaries = summaries_for(checked_settings({**defaults,
                                                         'peak_warning_enabled': '0',
                                                         'peak_warning_threshold': ''}))
            self.assertTrue(all(not row[-1] for row in summaries.values()))
            with self.assertRaisesRegex(ValueError, 'positive whole-number'):
                checked_settings({**defaults, 'peak_warning_threshold': ''})

    def test_optional_only_without_builtin_reference_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            settings = default_settings()
            for _, key in GUI_MARKERS:
                for field in ('marker_' + key + '_name', 'tb' + key + '1', 'tb' + key + '2'):
                    settings[field] = ''
            for index in range(1, 5):
                settings.update({f'custom{index}_name': f'M{index}',
                                 f'custom{index}_low': '-0.3',
                                 f'custom{index}_high': '+0.3'})
            settings = checked_settings(settings)
            book = Workbook()
            book.active.append(['Species/Genus', 'Order', 'M4', 'M1', 'M2', 'M3'])
            book.active.append(['Selected', 'Rodentia', 4000, 1000, 2000, 3000])
            book.active.append(['Other', 'Rodentia', 4100, 1100, 2100, 3100])
            book.save(root / 'ref.xlsx')
            (root / 'sample.tab').write_text('1000.09\t1\n2000.09\t1\n3000.09\t1\n4000.09\t1\n')
            output = root / 'results.csv'
            process(root, output, settings, reference_file=root / 'ref.xlsx')
            with output.open(encoding='utf-8-sig', newline='') as stream:
                records = list(csv.reader(stream))
            self.assertEqual(records[1][:2], ['sample.tab', '4'])
            self.assertEqual(records[1][records[0].index('M4')], '4000.1')
            self.assertEqual([row[records[0].index('Species/Genus')]
                              for row in records[2:-1]], ['Selected'])
            self.assertEqual(records[2][records[0].index('Order')], 'Rodentia')
            with self.assertRaisesRegex(ValueError, 'all twelve'):
                checked_settings({**settings, 'marker_P1_name': 'P1'})
            with self.assertRaisesRegex(ValueError, 'all twelve'):
                process(root, output, {**settings, 'marker_P1_name': 'P1'},
                        reference_file=root / 'ref.xlsx')
            empty = dict(settings)
            for index in range(1, 5):
                empty[f'custom{index}_name'] = ''
            with self.assertRaisesRegex(ValueError, 'at least one optional marker'):
                process(root, output, empty, reference_file=root / 'ref.xlsx')

    def test_clear_button_only_clears_builtin_fields(self):
        class Field:
            def __init__(self, value):
                self.value = value

            def set(self, value):
                self.value = value

        settings = default_settings()
        settings['custom1_name'] = 'Other marker'
        fields = {key: Field(value) for key, value in settings.items()}
        App.clear_builtin_markers(type('Form', (), {'vars': fields})())
        for _, key in GUI_MARKERS:
            for name in ('marker_' + key + '_name', 'tb' + key + '1', 'tb' + key + '2'):
                self.assertEqual(fields[name].value, '')
        self.assertEqual(fields['custom1_name'].value, 'Other marker')
        self.assertEqual(fields['tbminrows'].value, '4')
        self.assertEqual(fields['peak_warning_threshold'].value, '500')

    def test_peak_warning_entry_disabled_when_not_selected(self):
        class Entry:
            def configure(self, **options):
                self.state = options['state']

        class Variable:
            def __init__(self):
                self.value = '0'

            def get(self):
                return self.value

        variable = Variable()
        form = type('Form', (), {'vars': {'peak_warning_enabled': variable},
                                 'peak_warning_entry': Entry()})()
        App.update_peak_warning_field(form)
        self.assertEqual(form.peak_warning_entry.state, 'disabled')
        variable.value = '1'
        App.update_peak_warning_field(form)
        self.assertEqual(form.peak_warning_entry.state, 'normal')

    def test_defaults_load_until_settings_explicitly_saved(self):
        defaults = checked_settings(default_settings())
        self.assertEqual(defaults['tbminrows'], '4')
        self.assertEqual(defaults['remember_settings'], '0')
        for label, key in (('P1', 'P1'), ('A', 'A'), ("A'", 'A1'), ('E', 'E'),
                           ('G', 'G'), ("G'", 'G1')):
            self.assertEqual(defaults[f'tb{key}2'], '+1.3', label)
        with self.assertRaisesRegex(ValueError, 'one decimal place'):
            checked_settings({**defaults, 'tbB2': '+0.35'})
        with tempfile.TemporaryDirectory() as directory:
            legacy = Path(directory) / 'settings.json'
            legacy.write_text('{"indir":"old-path", "tbminrows":"7", "tbB2":"0.61"}')
            with patch('quickid.SETTINGS', legacy):
                loaded = load_settings()
            self.assertEqual(loaded['indir'], '')
            self.assertEqual(loaded['tbminrows'], '4')
            self.assertEqual(loaded['tbB2'], '+0.3')
            legacy.write_text('{"_defaults_version":2, "indir":"old-path", "tbB2":"+0.2", '
                              '"tbA12":"+0.3", "tbG12":"+0.3"}')
            with patch('quickid.SETTINGS', legacy):
                self.assertEqual(load_settings()['tbB2'], '+0.3')
                customized = {**defaults, 'indir': 'new-path', 'tbA12': '+0.8',
                              'marker_A1_name': 'Alpha prime'}
                save_settings(customized)
                loaded = load_settings()
                self.assertEqual(loaded['indir'], 'new-path')
                self.assertEqual(loaded['tbA12'], '+0.8')
                self.assertEqual(loaded['marker_A1_name'], 'Alpha prime')
                self.assertEqual(loaded['remember_settings'], '1')
                stop_remembering_settings()
                self.assertEqual(load_settings()['tbA12'], '+1.3')
                self.assertEqual(load_settings()['indir'], '')
                self.assertIn('"tbA12": "+0.8"', legacy.read_text())

    def test_run_does_not_save_unless_keep_settings_selected(self):
        values = {**default_settings(), 'outdir': 'result.csv'}
        form = type('Form', (), {'log': Mock(), 'values': lambda self: values,
                                 'report': lambda self, message: None})()
        with (patch('quickid.process', return_value=1),
              patch('quickid.save_settings') as save,
              patch('quickid.messagebox.showinfo')):
            App.run(form)
            save.assert_not_called()
            values['remember_settings'] = '1'
            App.run(form)
            save.assert_called_once()

    def test_paired_marker_upper_tolerance(self):
        row = dict.fromkeys(HEADINGS, '')
        row.update({'Easy Label': 'Taxon', 'A': '2000.0', "A'": '2100.0',
                    'G': '3000.0', "G'": '3100.0'})
        candidates, hits, _ = match([row], [Decimal('2100.9'), Decimal('3100.9')],
                                     default_settings())
        self.assertEqual(len(candidates), 1)
        self.assertEqual(hits[2], '2100.9')
        self.assertEqual(hits[11], '3100.9')

    def test_renamed_fixed_markers_select_reference_columns_and_keep_slot_tolerances(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            settings = default_settings()
            settings.update({'marker_P1_name': 'Q1', 'marker_A1_name': 'Alpha prime',
                             'marker_G1_name': 'Gamma prime', 'custom1_name': 'P1',
                             'custom1_low': '-0.3', 'custom1_high': '+0.3'})
            settings = checked_settings(settings)
            headings = ['Extra', 'Gamma prime', 'Q1', 'Alpha prime', 'P1',
                        *[name for name in HEADINGS if name not in ('P1', "A'", "G'")]]
            row = {name: '' for name in headings}
            row.update({'Extra': 'preserved', 'Species/Genus': 'Renamed taxon',
                        'Q1': '1000.0', 'Alpha prime': '2100.0',
                        'Gamma prime': '3100.0', 'P1': '4000.0'})
            book = Workbook()
            book.active.append(headings)
            book.active.append([row[name] for name in headings])
            book.save(root / 'ref.xlsx')
            (root / 'sample.txt').write_text('1000.1\t5\n2100.9\t5\n3100.9\t5\n4000.1\t5\n')
            output = root / 'out.csv'
            process(root, output, settings, reference_file=root / 'ref.xlsx')
            with output.open(encoding='utf-8-sig', newline='') as stream:
                records = list(csv.reader(stream))
            self.assertEqual(records[1][:2], ['sample.txt', '4'])
            for name, observed in [('Q1', '1000.1'), ('Alpha prime', '2100.9'),
                                   ('Gamma prime', '3100.9'), ('P1', '4000.1')]:
                self.assertEqual(records[1][records[0].index(name)], observed)
            self.assertEqual(records[2][records[0].index('Extra')], 'preserved')
            self.assertEqual(records[2][:2], ['', ''])
            with self.assertRaisesRegex(ValueError, "A'"):
                process(root, output, default_settings(), reference_file=root / 'ref.xlsx')
            with self.assertRaisesRegex(ValueError, 'unique'):
                checked_settings({**settings, 'marker_G1_name': 'q1'})

    def test_shuffled_reference_and_full_marker_match(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = Workbook()
            columns = list(reversed(HEADINGS))
            workbook.active.append(columns)
            masses = {name: str(1100 + index * 100) for index, name in enumerate(HEADINGS[3:])}
            taxon = {'Easy Label': 'Sample', 'Common Name': 'Animal', 'Species/Genus': 'Taxon', **masses}
            workbook.active.append([taxon[name] for name in columns])
            workbook.save(root / 'ref.xlsx')
            with (root / 'spec.txt').open('w') as stream:
                stream.write('m/z\tintensity\n')
                stream.writelines(f'{value}.09\t10\n' for value in masses.values())
            output = root / 'out.csv'
            processed = process(root, output, default_settings(), reference_file=root / 'ref.xlsx')
            self.assertEqual(processed, 1)
            with output.open(encoding='utf-8-sig', newline='') as stream:
                records = list(csv.reader(stream))
            self.assertEqual(records[0][:len(HEADINGS) + 2],
                             ['Sample name', 'QuickID marker count', *columns])
            self.assertEqual(records[1][0], 'spec.txt')
            self.assertEqual(records[1][1], '12')
            self.assertEqual(records[1][-1], '')
            self.assertEqual(records[2][:2], ['', ''])
            self.assertEqual(records[2][2:len(HEADINGS) + 2], [taxon[name] for name in columns])

    def test_duplicate_header_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'ref.xlsx'
            book = Workbook()
            book.active.append([*HEADINGS, 'p1'])
            book.active.append(['species', *([''] * len(HEADINGS))])
            book.save(path)
            with self.assertRaisesRegex(ValueError, 'P1'):
                read_reference(path)

    def test_second_member_of_marker_pair_can_select_candidate(self):
        first = dict.fromkeys(HEADINGS, '')
        second = dict.fromkeys(HEADINGS, '')
        first.update({'Easy Label': 'selected', 'A': '2000', "A'": '2100'})
        second.update({'Easy Label': 'other', 'A': '2200', "A'": '2300'})
        candidates, hits, _ = match([first, second], [Decimal('2100.0')], default_settings())
        self.assertEqual([row['Easy Label'] for row in candidates], ['selected'])
        self.assertEqual(hits[2], '2100.0')

    def test_tab_file_eight_optional_markers_and_minimum_four(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            book = Workbook()
            extras = [f'X{index}' for index in range(1, 9)]
            columns = list(reversed(HEADINGS + tuple(extras)))
            book.active.append(columns)
            taxon = {name: '' for name in HEADINGS}
            taxon.update({'Easy Label': 'Test', 'Common Name': 'Animal', 'Species/Genus': 'Taxon'})
            four = ['P1', 'A', "A'", 'B']
            taxon.update({name: str(1000 + index * 100) for index, name in enumerate(four)})
            taxon.update({name: str(3000 + index * 100) for index, name in enumerate(extras)})
            book.active.append([taxon.get(name, '') for name in columns])
            book.save(root / 'ref.xlsx')
            (root / 'spec.tab').write_text('m/z\tintensity\n' + ''.join(
                f'{taxon[name]}.09\t10\n' for name in four), encoding='utf-8')
            (root / 'ignore.csv').write_text('ignore', encoding='utf-8')
            output = root / 'result.csv'
            settings = checked_settings(default_settings())
            self.assertEqual(settings['tbminrows'], '4')
            process(root, output, settings, reference_file=root / 'ref.xlsx')
            with output.open(encoding='utf-8-sig', newline='') as stream:
                records = list(csv.reader(stream))
            self.assertEqual(records[1][1], '4')
            self.assertEqual(records[2][:2], ['', ''])
            self.assertEqual(records[2][columns.index('Easy Label') + 2], 'Test')  # Exactly four is sufficient.

            for index, name in enumerate(extras, start=1):
                settings.update({f'custom{index}_name': name, f'custom{index}_low': '-0.3',
                                 f'custom{index}_high': '+0.3'})
            (root / 'spec.tab').write_text('m/z\tintensity\n' + ''.join(
                f'{taxon[name]}.09\t10\n' for name in four + extras), encoding='utf-8')
            process(root, output, checked_settings(settings), reference_file=root / 'ref.xlsx')
            with output.open(encoding='utf-8-sig', newline='') as stream:
                records = list(csv.reader(stream))
            self.assertEqual(records[0][2:len(columns) + 2], columns)
            self.assertEqual(records[1][1], '12')
            self.assertEqual(records[2][2:len(columns) + 2], [taxon.get(name, '') for name in columns])

            settings['custom8_name'] = 'Absent'
            notices = []
            process(root, output, settings, notices.append, root / 'ref.xlsx')
            with output.open(encoding='utf-8-sig', newline='') as stream:
                self.assertEqual(list(csv.reader(stream))[0][2:len(columns) + 2], columns)
            self.assertTrue(any('column missing' in message for message in notices))

    def test_blank_heading_nonmarker_columns_and_headerless_peak_list(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            headings = ['Species/Genus', 'Common Name', '', 'Family', 'Order', *HEADINGS[3:]]
            book = Workbook()
            book.active.append(headings)
            row = ['Taxon', 'Animal', 'keep me', 'Bovidae', 'Artiodactyla',
                   *[str(1100 + index * 100) for index in range(12)]]
            book.active.append(row)
            book.save(root / 'ref.xlsx')
            peaks_folder = root / 'peak_lists'
            peaks_folder.mkdir()
            (peaks_folder / 'spec.txt').write_text('1100.09\t10\n' + ''.join(
                f'{value}.09\t10\n' for value in row[6:]), encoding='utf-8')
            output = root / 'out.csv'
            process(peaks_folder, output, default_settings(), reference_file=root / 'ref.xlsx')
            with output.open(encoding='utf-8-sig', newline='') as stream:
                records = list(csv.reader(stream))
            self.assertEqual(records[0][2:len(headings) + 2], headings)
            self.assertEqual(records[1][0], 'spec.txt')
            self.assertEqual(records[1][7], '1100.1')  # First peak was rounded and retained.
            self.assertEqual(records[2][:2], ['', ''])
            self.assertEqual(records[2][2:len(row) + 2], row)  # All nonmarker columns survive.


if __name__ == '__main__':
    unittest.main()
